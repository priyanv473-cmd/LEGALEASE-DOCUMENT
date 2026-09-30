import streamlit as st
import requests
from docx import Document
from fpdf import FPDF

st.set_page_config(
    page_title="LegalEase",
    page_icon="⚖️",
    layout="wide"
)

st.markdown("""
<style>
.stApp {
    background-color: #e8edf5;
}

.block-container {
    max-width: 1150px;
    padding-top: 2rem;
    padding-bottom: 3rem;
}

.brand {
    background: linear-gradient(135deg, #182848, #304a75);
    padding: 30px;
    border-radius: 18px;
    margin-bottom: 25px;
}

.logo-box {
    width: 64px;
    height: 64px;
    background: #d7b56d;
    color: #182848;
    border-radius: 15px;
    display: flex;
    align-items: center;
    justify-content: center;
    font-size: 22px;
    font-weight: 800;
    margin-bottom: 12px;
}

.brand-title {
    color: white;
    font-size: 38px;
    font-weight: 800;
}

.brand-subtitle {
    color: #cbd6e8;
    font-size: 16px;
}

.info-card {
    background: #f8fafc;
    padding: 24px;
    border-radius: 16px;
    border: 1px solid #d4ddea;
    margin-bottom: 20px;
}

.document-paper {
    background: #fffdf7;
    color: #222222;
    border: 1px solid #d8cdb7;
    border-radius: 12px;
    padding: 35px;
    margin-top: 15px;
    margin-bottom: 25px;
    box-shadow: 0 8px 25px rgba(70,55,30,0.12);
    font-size: 16px;
    line-height: 1.8;
}

.document-heading {
    color: #182848;
    font-size: 24px;
    font-weight: 700;
    border-bottom: 2px solid #d7b56d;
    padding-bottom: 10px;
    margin-top: 25px;
}

.download-area {
    background: #dde7f3;
    padding: 22px;
    border-radius: 15px;
    margin-top: 20px;
}

.disclaimer {
    background: #fff4d6;
    border-left: 5px solid #d7a632;
    padding: 15px;
    border-radius: 8px;
    color: #5d4a20;
    margin-top: 25px;
}

div.stButton > button {
    background-color: #182848;
    color: white;
    border: none;
    border-radius: 10px;
    font-weight: 700;
}

[data-testid="stDownloadButton"] button {
    background-color: #304a75;
    color: white;
    border-radius: 10px;
    border: none;
    font-weight: 600;
}
</style>
""", unsafe_allow_html=True)


st.markdown("""
<div class="brand">
    <div class="logo-box">LE</div>
    <div class="brand-title">LegalEase</div>
    <div class="brand-subtitle">
        AI-Powered Legal Document Generator
    </div>
</div>
""", unsafe_allow_html=True)


st.markdown("""
<div class="info-card">
<h2>Document Information</h2>
<p>Enter the required details below to generate your legal document.</p>
</div>
""", unsafe_allow_html=True)


col1, col2 = st.columns(2)

with col1:

    document_type = st.selectbox(
        "Document Type",
        [
            "Non-Disclosure Agreement",
            "Employment Contract",
            "Lease Agreement",
            "Service Agreement"
        ]
    )

    party1_name = st.text_input(
        "First Party",
        placeholder="Enter first party name"
    )

    party2_name = st.text_input(
        "Second Party",
        placeholder="Enter second party name"
    )


with col2:

    purpose = st.text_area(
        "Purpose / Agreement Details",
        placeholder="Describe the purpose of the agreement",
        height=125
    )

    date = st.date_input(
        "Document Date"
    )


st.write("")


if st.button(
    "Generate Legal Document",
    type="primary",
    use_container_width=True
):

    request_data = {
        "document_type": document_type,
        "party1_name": party1_name.strip(),
        "party2_name": party2_name.strip(),
        "purpose": purpose.strip(),
        "date": str(date)
    }

    try:

        with st.spinner(
            "LegalEase is preparing your document..."
        ):

            response = requests.post(
                "http://127.0.0.1:8000/generate",
                json=request_data,
                timeout=120
            )

        if response.status_code == 200:

            result = response.json()

            document_text = result.get(
                "document",
                ""
            )

            if not document_text.strip():

                st.error(
                    "Gemini did not return a document."
                )

            else:

                st.success(
                    "Legal document generated successfully!"
                )

                st.markdown(
                    '<div class="document-heading">Generated Legal Document</div>',
                    unsafe_allow_html=True
                )

                formatted_text = document_text.replace(
                    "\n",
                    "<br>"
                )

                st.markdown(
                    f"""
                    <div class="document-paper">
                        {formatted_text}
                    </div>
                    """,
                    unsafe_allow_html=True
                )


                docx = Document()

                docx.add_heading(
                    "LegalEase",
                    0
                )

                docx.add_paragraph(
                    f"Document Type: {document_type}"
                )

                docx.add_paragraph(
                    f"First Party: {party1_name}"
                )

                docx.add_paragraph(
                    f"Second Party: {party2_name}"
                )

                docx.add_paragraph(
                    f"Date: {date}"
                )

                docx.add_paragraph("")

                for paragraph in document_text.split("\n"):

                    if paragraph.strip():

                        docx.add_paragraph(
                            paragraph.strip()
                        )

                docx.add_paragraph("")

                docx.add_paragraph(
                    "Disclaimer: This is a general legal document "
                    "draft and not legal advice."
                )

                docx_path = "LegalEase_Document.docx"

                docx.save(docx_path)


                pdf = FPDF()

                pdf.set_auto_page_break(
                    auto=True,
                    margin=15
                )

                pdf.add_page()

                pdf.set_font(
                    "Helvetica",
                    "B",
                    20
                )

                pdf.cell(
                    0,
                    12,
                    "LegalEase",
                    ln=True
                )

                pdf.set_font(
                    "Helvetica",
                    "",
                    11
                )

                pdf.cell(
                    0,
                    8,
                    f"Document Type: {document_type}",
                    ln=True
                )

                pdf.cell(
                    0,
                    8,
                    f"First Party: {party1_name}",
                    ln=True
                )

                pdf.cell(
                    0,
                    8,
                    f"Second Party: {party2_name}",
                    ln=True
                )

                pdf.cell(
                    0,
                    8,
                    f"Date: {date}",
                    ln=True
                )

                pdf.ln(8)

                safe_text = document_text.encode(
                    "latin-1",
                    "replace"
                ).decode(
                    "latin-1"
                )

                for paragraph in safe_text.split("\n"):

                    if paragraph.strip():

                        pdf.multi_cell(
                            0,
                            7,
                            paragraph.strip()
                        )

                        pdf.ln(2)

                pdf.ln(5)

                pdf.set_font(
                    "Helvetica",
                    "I",
                    9
                )

                pdf.multi_cell(
                    0,
                    6,
                    "Disclaimer: This is a general legal document "
                    "draft and not legal advice."
                )

                pdf_path = "LegalEase_Document.pdf"

                pdf.output(pdf_path)


                st.markdown(
                    """
                    <div class="download-area">
                        <h3>Download Your Document</h3>
                        <p>Choose your required format.</p>
                    </div>
                    """,
                    unsafe_allow_html=True
                )


                download_col1, download_col2 = st.columns(2)

                with download_col1:

                    with open(
                        docx_path,
                        "rb"
                    ) as file:

                        st.download_button(
                            "📘 Download DOCX",
                            file,
                            file_name="LegalEase_Document.docx",
                            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                            use_container_width=True
                        )


                with download_col2:

                    with open(
                        pdf_path,
                        "rb"
                    ) as file:

                        st.download_button(
                            "📕 Download PDF",
                            file,
                            file_name="LegalEase_Document.pdf",
                            mime="application/pdf",
                            use_container_width=True
                        )


        else:

            st.error(
                f"Document generation failed. Status code: {response.status_code}"
            )

            st.write(response.text)


    except requests.exceptions.ConnectionError:

        st.error(
            "Cannot connect to FastAPI backend. "
            "Please make sure uvicorn is running."
        )


    except requests.exceptions.Timeout:

        st.error(
            "The request took too long. Please try again."
        )


    except Exception as e:

        st.error(
            f"Something went wrong: {str(e)}"
        )


st.markdown("""
<div class="disclaimer">
<b>Disclaimer:</b> LegalEase provides general legal document
drafts for informational purposes only. It is not a substitute
for professional legal advice.
</div>
""", unsafe_allow_html=True)